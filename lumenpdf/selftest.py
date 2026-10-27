"""Autoteste executado dentro do programa empacotado (LumenPDF.exe --selftest saida.txt).
Garante que o instalador leva tudo o que é necessário: assinatura, OCR, conversões e operações."""
import datetime
import os
import tempfile
import traceback


def run(log_path=None):
    lines = []

    def log(msg):
        lines.append(msg)
        print(msg, flush=True)
    try:
        import pymupdf
        from cryptography import x509
        from cryptography.hazmat.primitives import hashes, serialization
        from cryptography.hazmat.primitives.asymmetric import rsa
        from cryptography.hazmat.primitives.serialization import pkcs12
        from cryptography.x509.oid import NameOID
        from . import ops, signing
        print_ = log

        tmp = tempfile.mkdtemp()
        doc = pymupdf.open()
        for i in range(3):
            p = doc.new_page()
            p.insert_text((72, 100), f"Página {i + 1} contrato CPF 123.456.789-00", fontsize=12)
        data = doc.tobytes()

        # assinatura digital com carimbo (logo)
        k = rsa.generate_private_key(public_exponent=65537, key_size=2048)
        n = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "TESTE CI:00000000000")])
        c = (x509.CertificateBuilder().subject_name(n).issuer_name(n).public_key(k.public_key()).serial_number(1)
             .not_valid_before(datetime.datetime(2025, 1, 1)).not_valid_after(datetime.datetime(2035, 1, 1)).sign(k, hashes.SHA256()))
        pfx = os.path.join(tmp, "t.pfx")
        open(pfx, "wb").write(pkcs12.serialize_key_and_certificates(b"t", k, c, None, serialization.BestAvailableEncryption(b"1")))
        signer = signing.load_pfx(pfx, "1")
        out = signing.sign_pdf_bytes(data, signer, 0, (300, 600, 540, 680), 842, "Teste", "Goiânia")
        res = signing.validate(out)
        assert res and res[0]["intact"] and res[0]["valid"], res
        # segunda assinatura (outro certificado) não pode invalidar a primeira
        k2 = rsa.generate_private_key(public_exponent=65537, key_size=2048)
        n2 = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "SEGUNDO CI")])
        c2 = (x509.CertificateBuilder().subject_name(n2).issuer_name(n2).public_key(k2.public_key()).serial_number(2)
              .not_valid_before(datetime.datetime(2025, 1, 1)).not_valid_after(datetime.datetime(2035, 1, 1)).sign(k2, hashes.SHA256()))
        pfx2 = os.path.join(tmp, "t2.pfx")
        open(pfx2, "wb").write(pkcs12.serialize_key_and_certificates(b"t2", k2, c2, None, serialization.BestAvailableEncryption(b"1")))
        out2 = signing.sign_pdf_bytes(out, signing.load_pfx(pfx2, "1"), 0, (50, 600, 290, 680), 842, "Teste 2", "")
        res2 = signing.validate(out2)
        assert len(res2) == 2 and all(r["intact"] and r["valid"] for r in res2), res2
        print_("assinatura ok (duas assinaturas válidas)")

        # leitura de código de barras e QR (boleto / Pix)
        import zxingcpp
        from . import pagamentos
        qr = zxingcpp.create_barcode("00020126580014br.gov.bcb.pix0136123e4567-e12b-12d1-a456-4266554400005204000053039865802BR5913Fulano de Tal6008BRASILIA62070503***63041D3D", zxingcpp.BarcodeFormat.QRCode).to_svg()
        qd = pymupdf.open("pdf", pymupdf.open(stream=qr.encode(), filetype="svg").convert_to_pdf())
        dq = pymupdf.open(); dq.new_page().show_pdf_page(pymupdf.Rect(50, 50, 250, 250), qd, 0)
        found = pagamentos.scan(dq)
        assert found and found[0].kind == "pix" and found[0].ok, found
        print_("boleto/Pix ok")

        # posição das páginas (texto de cabeça para baixo é detectado e corrigido)
        dr = pymupdf.open()
        pr = dr.new_page()
        for k in range(12):
            pr.insert_text((60, 90 + 18 * k), "Cláusula de cessão de créditos entre as partes contratantes", fontsize=12)
        pr.set_rotation(180)
        assert ops.orientation_fixes(dr) == {0: 0}
        assert ops.tessdata_dir(), "dados do OCR"
        print_("posição das páginas ok")

        # tradução offline: bibliotecas empacotadas (o pacote de idioma é baixado na primeira tradução)
        import ctranslate2
        import sentencepiece
        from . import traducao, artigos
        assert traducao.detect("The results of this study show that the animals were kept in poor conditions for a long time.") == "en"
        assert artigos.Article("Título", ["Maria da Silva"], 2020, "Revista").abnt().startswith("SILVA, Maria da.")
        print_(f"tradução ok (ctranslate2 {ctranslate2.__version__})")

        # conversões e operações
        ops.to_docx(data, os.path.join(tmp, "x.docx"))
        ops.to_xlsx(pymupdf.open("pdf", data), os.path.join(tmp, "x.xlsx"))
        ops.compress(pymupdf.open("pdf", data), "media")
        d2 = pymupdf.open("pdf", data)
        ops.watermark(d2, "TESTE")
        ops.page_numbers(d2)
        assert ops.redact_pattern(d2, ops.PATTERNS["CPF"]) == 3
        ops.apply_redactions(d2)
        ops.encrypted_bytes(d2, "1", "2", ["print"])
        print_("operações ok")

        # OCR com os idiomas embutidos (português)
        from lumenpdf import resource
        assert os.path.isfile(resource("assets", "tessdata", "por.traineddata")), "tessdata não embutido"
        src = pymupdf.open()
        pg = src.new_page()
        pg.insert_text((60, 120), "Declaro que RECEBIDO o valor da nota fiscal", fontsize=20)
        img_doc = pymupdf.open()
        ip = img_doc.new_page()
        ip.insert_image(ip.rect, pixmap=pg.get_pixmap(dpi=200))
        assert ops.ocr(img_doc, language="por", dpi=200) == 1
        assert "RECEBIDO" in img_doc[0].get_text(), img_doc[0].get_text()
        print_("ocr ok")

        # conversão de Office sem Office/LibreOffice
        from docx import Document
        from openpyxl import Workbook
        dx = os.path.join(tmp, "a.docx")
        d = Document()
        d.add_heading("Contrato de teste", 1)
        d.add_paragraph("Parágrafo com acentuação: ação, cessão.")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(1, 1).text = "Título", "R$ 10,00"
        d.save(dx)
        ops.builtin_office_to_pdf(dx, os.path.join(tmp, "a.pdf"))
        assert "cessão" in pymupdf.open(os.path.join(tmp, "a.pdf"))[0].get_text()
        xl = os.path.join(tmp, "b.xlsx")
        wb = Workbook()
        wb.active.append(["Título", "Valor"])
        wb.active.append(["NF 1", 1234.5])
        wb.save(xl)
        ops.builtin_office_to_pdf(xl, os.path.join(tmp, "b.pdf"))
        assert "1.234,50" in pymupdf.open(os.path.join(tmp, "b.pdf"))[0].get_text()
        print_("office embutido ok")
        # carimbo de tempo + LTV (só quando houver internet liberada, ex.: no build)
        if os.environ.get("LUMEN_NET_TEST") == "1":
            rep = {}
            out = signing.sign_pdf_bytes(data, signer, 0, (300, 600, 540, 680), 842, "TSA", "CI",
                                         tsa_url=signing.DEFAULT_TSA, ltv=True, report=rep)
            res = signing.validate(out)[0]
            assert rep["tsa"] and res["tsa_time"] and res["intact"], (rep, res)
            log(f"carimbo de tempo ok (LTV: {rep['ltv']})")

        # edição de texto mantendo a fonte
        dd = pymupdf.open("pdf", data)
        sp = ops.span_at(dd[0], (80, 95))
        used = ops.replace_span(dd[0], sp, "Texto editado ção")
        assert "editado" in dd[0].get_text(), dd[0].get_text()
        log(f"edição de texto ok ({used})")

        # gerenciador de impressão: imposição e saída via QPdfWriter (mesmo caminho da impressora)
        from PySide6.QtWidgets import QApplication
        _app = QApplication.instance() or QApplication([])
        from PySide6.QtGui import QPdfWriter, QPainter, QPageSize
        from . import print_manager as pmod
        src = pmod.Source("teste.pdf", pymupdf.open("pdf", data))
        for mode in ("size", "multiple", "booklet", "poster"):
            o = dict(pmod.DEFAULT_OPTS, mode=mode, footer=True, watermark="CÓPIA", gray=(mode == "multiple"))
            sheets = pmod.impose([src], o, 559, 783)
            out_pdf = os.path.join(tmp, f"print_{mode}.pdf")
            wtr = QPdfWriter(out_pdf)
            wtr.setPageSize(QPageSize(QPageSize.A4))
            wtr.setResolution(150)
            pnt = QPainter(wtr)
            kk = wtr.resolution() / 72
            for i, sh in enumerate(sheets):
                if i:
                    wtr.newPage()
                pmod.draw_sheet(pnt, sh, [src], o, kk, (0, 0), 559, 783, sheet_no=i + 1, total=len(sheets))
            pnt.end()
            assert pymupdf.open(out_pdf).page_count == len(sheets), mode
        log("impressão ok (tamanho, várias, livreto, pôster)")

        # módulos carregados sob demanda pela interface
        import PySide6.QtPrintSupport  # noqa
        import PySide6.QtSvg  # noqa
        import PySide6.QtNetwork  # noqa
        try:
            import PySide6.QtTextToSpeech  # noqa
            log("voz ok")
        except ImportError as e:
            log(f"AVISO: leitura em voz alta indisponível ({e})")
        from . import wincert, updater, system, errors  # noqa
        log("SELFTEST OK")
        ok = True
    except Exception:
        log("SELFTEST FALHOU\n" + traceback.format_exc())
        ok = False
    if log_path:
        with open(log_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
    return 0 if ok else 1
